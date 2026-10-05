"""Load an Instamart Sales report xlsx into the seller tables.

    python -m scripts.load_instamart_sales <tenant_id> file1.xlsx [file2.xlsx ...]

Parses the report's two sheets and upserts them by `upsert_key`
(instamart_seller_store_daily, instamart_brand_city_daily). Re-running the same
file overwrites in place — safe to load overlapping windows. `scrape_job_id`
is left null for a manual load.

Same ON CONFLICT (upsert_key) DO UPDATE idiom as the Zepto seller storage.
"""
import asyncio
import datetime as dt
import sys
import uuid

import openpyxl
from sqlalchemy.dialects.postgresql import insert

from app.core.database import AsyncSessionLocal
from app.models import InstamartBrandCityDaily, InstamartSellerStoreDaily
from app.utils.logger import logger
from app.utils.time import now_ist


def _clean(v):
    if isinstance(v, str):
        return v.replace("\xa0", " ").strip() or None
    return v


def _date(v):
    return v.date() if isinstance(v, dt.datetime) else v


def _int(v):
    return None if v in (None, "") else int(float(v))


def _float(v):
    return None if v in (None, "") else float(v)


def parse(path: str) -> tuple[list[dict], list[dict]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    store, brand = [], []

    ws = wb["Sales Report"]
    it = ws.iter_rows(values_only=True)
    col = {name: i for i, name in enumerate(next(it))}

    def g(r, name):
        i = col.get(name)
        return r[i] if i is not None and i < len(r) else None

    for r in it:
        if g(r, "STORE_ID") is None or g(r, "ITEM_CODE") is None:
            continue
        store.append({
            "date": _date(g(r, "ORDERED_DATE")),
            "city": _clean(g(r, "CITY")),
            "area_name": _clean(g(r, "AREA_NAME")),
            "store_id": str(g(r, "STORE_ID")),
            "l1_category": _clean(g(r, "L1_CATEGORY")),
            "l2_category": _clean(g(r, "L2_CATEGORY")),
            "l3_category": _clean(g(r, "L3_CATEGORY")),
            "product_name": _clean(g(r, "PRODUCT_NAME")),
            "variant": _clean(g(r, "VARIANT")),
            "item_code": str(g(r, "ITEM_CODE")),
            "is_combo": str(g(r, "COMBO")).strip().lower() == "yes",
            "combo_item_code": _clean(g(r, "COMBO_ITEM_CODE")),
            "combo_units_sold": _int(g(r, "COMBO_UNITS_SOLD")) or 0,
            "base_mrp": _float(g(r, "BASE_MRP")),
            "units_sold": _int(g(r, "UNITS_SOLD")) or 0,
            "gmv": _float(g(r, "GMV")) or 0.0,
        })

    if "Brand Metrics" in wb.sheetnames:
        rows = list(wb["Brand Metrics"].iter_rows(values_only=True))
        col = {name: i for i, name in enumerate(rows[1])}       # row 0 = note
        for r in rows[2:]:
            city = _clean(r[col["CITY"]])
            if not city or city == "CITY":
                continue
            brand.append({
                "date": _date(r[col["DATE"]]),
                "city": city,
                "ntb_buyers": _int(r[col["NTB_BUYERS"]]),
                "brand_impressions": _int(r[col["BRAND_IMPRESSIONS"]]),
                "brand_gmv": _float(r[col["BRAND_GMV"]]),
                "brand_orders": _int(r[col["BRAND_ORDERS"]]),
            })
    return store, brand


def _key(*parts) -> str:
    return "|".join(str(p) for p in parts)


async def _upsert(session, model, rows: list[dict]) -> int:
    if not rows:
        return 0
    deduped = {r["upsert_key"]: r for r in rows}
    rows = list(deduped.values())
    update_cols = [c.name for c in model.__table__.columns
                   if c.name not in {"id", "upsert_key"}]
    cols = max(1, len(model.__table__.columns))
    chunk = max(1, 32000 // cols)
    for i in range(0, len(rows), chunk):
        stmt = (
            insert(model)
            .values(rows[i:i + chunk])
            .on_conflict_do_update(
                index_elements=["upsert_key"],
                set_={c: insert(model).excluded[c] for c in update_cols},
            )
        )
        await session.execute(stmt)
    return len(rows)


async def load(tenant_id: str, paths: list[str]) -> None:
    tid = uuid.UUID(str(tenant_id))
    now = now_ist()
    async with AsyncSessionLocal() as session:
        for path in paths:
            store, brand = parse(path)
            for r in store:
                r["tenant_id"] = tid
                r["platform"] = "instamart"
                r["scrape_job_id"] = None
                r["scraped_at"] = now
                r["upsert_key"] = _key("instamart", tid, r["date"], r["store_id"],
                                       r["item_code"])
            for r in brand:
                r["tenant_id"] = tid
                r["platform"] = "instamart"
                r["scrape_job_id"] = None
                r["scraped_at"] = now
                r["upsert_key"] = _key("instamart", tid, r["date"], r["city"])

            n1 = await _upsert(session, InstamartSellerStoreDaily, store)
            n2 = await _upsert(session, InstamartBrandCityDaily, brand)
            await session.commit()
            logger.info(f"{path.split(chr(92))[-1].split('/')[-1]}: "
                        f"{n1} store rows, {n2} brand rows upserted")


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("usage: python -m scripts.load_instamart_sales <tenant_id> file.xlsx ...")
        raise SystemExit(2)
    asyncio.run(load(sys.argv[1], sys.argv[2:]))
