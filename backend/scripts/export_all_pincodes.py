"""
Export all city-wise pincodes from endpoints.py to Excel.
Run from backend/:
    python -m scripts.export_all_pincodes
Output: zepto_all_city_pincodes.xlsx
"""
import sys
from pathlib import Path
from datetime import datetime

sys.path.insert(0, str(Path(__file__).parent.parent))

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from scraper.platforms.zepto.dark_store.endpoints import (
    DELHI_NCR_PINCODES,
    BANGALORE_PINCODES,
    CHENNAI_PINCODES,
    HYDERABAD_PINCODES,
    KOLKATA_PINCODES,
    PUNE_PINCODES,
    MUMBAI_PINCODES,
    SURAT_PINCODES,
    JAIPUR_PINCODES,
    AHMEDABAD_PINCODES,
    COIMBATORE_PINCODES,
    NASHIK_PINCODES,
    MYSURU_PINCODES,
    INDORE_PINCODES,
    DEHRADUN_PINCODES,
    NAGPUR_PINCODES,
    KANPUR_PINCODES,
    LUCKNOW_PINCODES,
    KOCHI_PINCODES,
    LUDHIANA_PINCODES,
    JALANDHAR_PINCODES,
    PRAYAGRAJ_PINCODES,
    VIJAYAWADA_PINCODES,
    VADODARA_PINCODES,
    AGRA_PINCODES,
    BAREILLY_PINCODES,
    CHANDIGARH_PINCODES,
    MADURAI_PINCODES,
    MOHALI_PINCODES,
    UDAIPUR_PINCODES,
    MEERUT_PINCODES,
    AMRITSAR_PINCODES,
    GORAKHPUR_PINCODES,
    HUBBALLI_PINCODES,
    KARIMNAGAR_PINCODES,
    KOTA_PINCODES,
    PALAKKAD_PINCODES,
    PATIALA_PINCODES,
    VARANASI_PINCODES,
    WARANGAL_PINCODES,
    AMBALA_PINCODES,
    DAVANGERE_PINCODES,
    GUNTUR_PINCODES,
    HARIDWAR_PINCODES,
    KARNAL_PINCODES,
    PANIPAT_PINCODES,
    KURUKSHETRA_PINCODES,
    PUDUCHERRY_PINCODES,
    RAJKOT_PINCODES,
    VELLORE_PINCODES,
    SONIPAT_PINCODES,
    BELAGAVI_PINCODES,
    BHIWADI_PINCODES,
    HAPUR_PINCODES,
    HISAR_PINCODES,
    HOSUR_PINCODES,
    MEHSANA_PINCODES,
    PANCHKULA_PINCODES,
    REWARI_PINCODES,
    SAHARANPUR_PINCODES,
    VALSAD_PINCODES,
    TUMKURU_PINCODES,
    CSN_PINCODES,
    MANGALURU_PINCODES,
    THIRUVANANTHAPURAM_PINCODES,
    BHUBANESWAR_PINCODES,
    GUWAHATI_PINCODES,
    TRICHY_PINCODES,
    SALEM_PINCODES,
)

ALL_CITIES = {
    "Delhi NCR": DELHI_NCR_PINCODES,
    "Bangalore": BANGALORE_PINCODES,
    "Chennai": CHENNAI_PINCODES,
    "Hyderabad": HYDERABAD_PINCODES,
    "Kolkata": KOLKATA_PINCODES,
    "Pune": PUNE_PINCODES,
    "Mumbai": MUMBAI_PINCODES,
    "Surat": SURAT_PINCODES,
    "Jaipur": JAIPUR_PINCODES,
    "Ahmedabad": AHMEDABAD_PINCODES,
    "Coimbatore": COIMBATORE_PINCODES,
    "Nashik": NASHIK_PINCODES,
    "Mysuru": MYSURU_PINCODES,
    "Indore": INDORE_PINCODES,
    "Dehradun": DEHRADUN_PINCODES,
    "Nagpur": NAGPUR_PINCODES,
    "Kanpur": KANPUR_PINCODES,
    "Lucknow": LUCKNOW_PINCODES,
    "Kochi": KOCHI_PINCODES,
    "Ludhiana": LUDHIANA_PINCODES,
    "Jalandhar": JALANDHAR_PINCODES,
    "Prayagraj": PRAYAGRAJ_PINCODES,
    "Vijayawada": VIJAYAWADA_PINCODES,
    "Vadodara": VADODARA_PINCODES,
    "Agra": AGRA_PINCODES,
    "Bareilly": BAREILLY_PINCODES,
    "Chandigarh": CHANDIGARH_PINCODES,
    "Madurai": MADURAI_PINCODES,
    "Mohali": MOHALI_PINCODES,
    "Udaipur": UDAIPUR_PINCODES,
    "Meerut": MEERUT_PINCODES,
    "Amritsar": AMRITSAR_PINCODES,
    "Gorakhpur": GORAKHPUR_PINCODES,
    "Hubballi": HUBBALLI_PINCODES,
    "Karimnagar": KARIMNAGAR_PINCODES,
    "Kota": KOTA_PINCODES,
    "Palakkad": PALAKKAD_PINCODES,
    "Patiala": PATIALA_PINCODES,
    "Varanasi": VARANASI_PINCODES,
    "Warangal": WARANGAL_PINCODES,
    "Ambala": AMBALA_PINCODES,
    "Davangere": DAVANGERE_PINCODES,
    "Guntur": GUNTUR_PINCODES,
    "Haridwar": HARIDWAR_PINCODES,
    "Karnal": KARNAL_PINCODES,
    "Panipat": PANIPAT_PINCODES,
    "Kurukshetra": KURUKSHETRA_PINCODES,
    "Puducherry": PUDUCHERRY_PINCODES,
    "Rajkot": RAJKOT_PINCODES,
    "Vellore": VELLORE_PINCODES,
    "Sonipat": SONIPAT_PINCODES,
    "Belagavi": BELAGAVI_PINCODES,
    "Bhiwadi": BHIWADI_PINCODES,
    "Hapur": HAPUR_PINCODES,
    "Hisar": HISAR_PINCODES,
    "Hosur": HOSUR_PINCODES,
    "Mehsana": MEHSANA_PINCODES,
    "Panchkula": PANCHKULA_PINCODES,
    "Rewari": REWARI_PINCODES,
    "Saharanpur": SAHARANPUR_PINCODES,
    "Valsad": VALSAD_PINCODES,
    "Tumkuru": TUMKURU_PINCODES,
    "Chhatrapati Sambhaji Nagar": CSN_PINCODES,
    "Mangaluru": MANGALURU_PINCODES,
    "Thiruvananthapuram": THIRUVANANTHAPURAM_PINCODES,
    "Bhubaneswar": BHUBANESWAR_PINCODES,
    "Guwahati": GUWAHATI_PINCODES,
    "Trichy": TRICHY_PINCODES,
    "Salem": SALEM_PINCODES,
}

HDR_FONT   = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
HDR_FILL   = PatternFill("solid", fgColor="2E4057")
HDR_ALIGN  = Alignment(horizontal="center", vertical="center")
ALT_FILL   = PatternFill("solid", fgColor="F0F4F8")
CITY_FILL  = PatternFill("solid", fgColor="048A81")
CITY_FONT  = Font(name="Calibri", size=10, bold=True, color="FFFFFF")


def col_width(ws, col_idx, min_w=12, max_w=40):
    max_len = 0
    for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
        for cell in row:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
    ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, min_w), max_w)


def main():
    wb = openpyxl.Workbook()

    # ── Sheet 1: all pincodes flat ──────────────────────────────────────────
    ws_all = wb.active
    ws_all.title = "All Pincodes"

    headers = ["City", "Pincode", "Area Label"]
    ws_all.append(headers)
    for cell in ws_all[1]:
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.alignment = HDR_ALIGN
    ws_all.row_dimensions[1].height = 18

    row_num = 2
    for city, pincodes in ALL_CITIES.items():
        for pincode, label in pincodes.items():
            ws_all.cell(row=row_num, column=1, value=city)
            ws_all.cell(row=row_num, column=2, value=pincode)
            ws_all.cell(row=row_num, column=3, value=label)
            if row_num % 2 == 0:
                for col in range(1, 4):
                    ws_all.cell(row=row_num, column=col).fill = ALT_FILL
            row_num += 1

    for i in range(1, 4):
        col_width(ws_all, i)
    ws_all.freeze_panes = "A2"
    ws_all.auto_filter.ref = ws_all.dimensions

    # ── Sheet 2: summary (city + pincode count) ─────────────────────────────
    ws_sum = wb.create_sheet("Summary")
    ws_sum.append(["City", "Pincode Count"])
    for cell in ws_sum[1]:
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.alignment = HDR_ALIGN
    ws_sum.row_dimensions[1].height = 18

    for i, (city, pincodes) in enumerate(ALL_CITIES.items(), start=2):
        ws_sum.cell(row=i, column=1, value=city)
        ws_sum.cell(row=i, column=2, value=len(pincodes))
        if i % 2 == 0:
            for col in range(1, 3):
                ws_sum.cell(row=i, column=col).fill = ALT_FILL

    total_row = len(ALL_CITIES) + 2
    ws_sum.cell(row=total_row, column=1, value="TOTAL")
    ws_sum.cell(row=total_row, column=2, value=sum(len(p) for p in ALL_CITIES.values()))
    for col in range(1, 3):
        ws_sum.cell(row=total_row, column=col).font = CITY_FONT
        ws_sum.cell(row=total_row, column=col).fill = CITY_FILL

    for i in range(1, 3):
        col_width(ws_sum, i)
    ws_sum.freeze_panes = "A2"

    # ── Sheet 3: one sheet per city ─────────────────────────────────────────
    for city, pincodes in ALL_CITIES.items():
        safe = city[:28].replace("/", "-").replace(":", "")
        ws = wb.create_sheet(safe)
        ws.append(["Pincode", "Area Label"])
        for cell in ws[1]:
            cell.font = HDR_FONT
            cell.fill = HDR_FILL
            cell.alignment = HDR_ALIGN
        ws.row_dimensions[1].height = 18
        for r, (pincode, label) in enumerate(pincodes.items(), start=2):
            ws.cell(row=r, column=1, value=pincode)
            ws.cell(row=r, column=2, value=label)
            if r % 2 == 0:
                for col in range(1, 3):
                    ws.cell(row=r, column=col).fill = ALT_FILL
        for i in range(1, 3):
            col_width(ws, i)
        ws.freeze_panes = "A2"

    out = Path(__file__).parent.parent / "zepto_all_city_pincodes.xlsx"
    wb.save(out)
    print(f"\nCities: {len(ALL_CITIES)}")
    print(f"Total pincodes: {sum(len(p) for p in ALL_CITIES.values())}")
    print(f"Saved → {out}")

    for city, pincodes in ALL_CITIES.items():
        print(f"  {city}: {len(pincodes)}")


if __name__ == "__main__":
    main()
