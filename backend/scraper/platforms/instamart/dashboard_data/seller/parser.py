"""Turn a downloaded Sales report xlsx into rows for the two seller tables.

The workbook has two sheets at two different grains, and they must stay apart:

  "Sales Report"  one row per ORDERED_DATE x STORE_ID x ITEM_CODE, only where
                  something sold. Column A is a 0-based row index, not data.
  "Brand Metrics" a note row, then a header on row 2, then one row per
                  DATE x CITY for the whole brand. Numbers arrive as strings and
                  are blank for a city with impressions but no sales.

Never sum one against the other: they hold the same rupees at different
resolutions (the same rule as zepto_seller_sales vs
zepto_seller_product_city_daily).

Verified against the real files on 2026-09-22: 6,890 rows across 1-21 Sep,
GMV == BASE_MRP * UNITS_SOLD on every row, and (date, store_id, item_code) is
unique — which is what makes it a safe upsert key.
"""
import datetime as dt
from pathlib import Path

import openpyxl

SALES_SHEET = "Sales Report"
BRAND_SHEET = "Brand Metrics"


def _clean(value):
    """Trim, and fold the non-breaking space Instamart puts in some category
    names ("fresh\xa0bakery") so joins and GROUP BYs behave."""
    if isinstance(value, str):
        return value.replace("\xa0", " ").strip() or None
    return value


def _date(value):
    return value.date() if isinstance(value, dt.datetime) else value


def _int(value):
    return None if value in (None, "") else int(float(value))


def _float(value):
    return None if value in (None, "") else float(value)


def parse_store_daily(path: str | Path) -> list[dict]:
    """One dict per (date, store, item) — the "Sales Report" sheet."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if SALES_SHEET not in wb.sheetnames:
        raise ValueError(f"{path}: no {SALES_SHEET!r} sheet — is this a sales report?")

    rows_iter = wb[SALES_SHEET].iter_rows(values_only=True)
    header = next(rows_iter)
    col = {name: i for i, name in enumerate(header)}

    def cell(row, name):
        i = col.get(name)
        return row[i] if i is not None and i < len(row) else None

    out: list[dict] = []
    for row in rows_iter:
        # A report for a day with no sales still carries its header row.
        if cell(row, "STORE_ID") is None or cell(row, "ITEM_CODE") is None:
            continue
        out.append({
            "date": _date(cell(row, "ORDERED_DATE")),
            "city": _clean(cell(row, "CITY")),
            "area_name": _clean(cell(row, "AREA_NAME")),
            # Instamart's podId — the same id the public scraper stores as
            # search_listings.merchant_id, which is what lets private sales
            # join to public price/stock per store.
            "store_id": str(cell(row, "STORE_ID")),
            "l1_category": _clean(cell(row, "L1_CATEGORY")),
            "l2_category": _clean(cell(row, "L2_CATEGORY")),
            "l3_category": _clean(cell(row, "L3_CATEGORY")),
            "product_name": _clean(cell(row, "PRODUCT_NAME")),
            "variant": _clean(cell(row, "VARIANT")),
            "item_code": str(cell(row, "ITEM_CODE")),
            "is_combo": str(cell(row, "COMBO")).strip().lower() == "yes",
            "combo_item_code": _clean(cell(row, "COMBO_ITEM_CODE")),
            "combo_units_sold": _int(cell(row, "COMBO_UNITS_SOLD")) or 0,
            "base_mrp": _float(cell(row, "BASE_MRP")),
            "units_sold": _int(cell(row, "UNITS_SOLD")) or 0,
            "gmv": _float(cell(row, "GMV")) or 0.0,
        })
    return out


def parse_brand_city(path: str | Path) -> list[dict]:
    """One dict per (date, city) — the "Brand Metrics" sheet, if present."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if BRAND_SHEET not in wb.sheetnames:
        return []

    rows = list(wb[BRAND_SHEET].iter_rows(values_only=True))
    if len(rows) < 3:
        return []
    col = {name: i for i, name in enumerate(rows[1])}   # row 0 is a note
    if "CITY" not in col:
        return []

    out: list[dict] = []
    for row in rows[2:]:
        city = _clean(row[col["CITY"]])
        if not city or city == "CITY":
            continue
        out.append({
            "date": _date(row[col["DATE"]]),
            "city": city,
            # Every metric is nullable: a city can show impressions with no sales.
            "ntb_buyers": _int(row[col["NTB_BUYERS"]]),
            "brand_impressions": _int(row[col["BRAND_IMPRESSIONS"]]),
            "brand_gmv": _float(row[col["BRAND_GMV"]]),
            "brand_orders": _int(row[col["BRAND_ORDERS"]]),
        })
    return out


def parse(path: str | Path) -> tuple[list[dict], list[dict]]:
    """Both sheets in one go: (store_daily rows, brand_city rows)."""
    return parse_store_daily(path), parse_brand_city(path)
