"""Blinkit seller-hub (seller.blinkit.com) — sales scrape.

Unlike scraper/platforms/blinkit/dashboard_data/seller/scraper.py (the OLD
partnersbiz.com dashboard), this CANNOT be a plain httpx client from OUTSIDE
the browser. This domain's Cloudflare bot management blocks a bare request
even with a live session's exact cookies/headers replayed — proven twice in
platform_auth/marketplaces/blinkit/seller_new.py (Phase B, 2026-09-30). So a
real browser has to open the page first, every time.

BUT once inside that authenticated page, an in-page fetch() to the sales
endpoints DOES work — corrected 2026-10-01 after two earlier false starts.
The first in-page fetch() attempt (against a DIFFERENT set of endpoints,
during auth development: app/config, app/feed, notifications, faq) failed
with "Missing required headers for AuthLevelRequired", and that finding was
wrongly generalized to "fetch never works on this domain." Re-tested
specifically against sales/performance/metrics with the real SPA's headers
captured live: four headers a hand-built fetch doesn't normally send are
missing — `access_token` (as a HEADER, not just the cookie), `app_client`,
`x-api-key`, `x-gr-seller-id`. Add those four (the first and last are read
live from cookies each run; the middle two are static values baked into the
SPA's own JS bundle) and the identical fetch returns 200 with real data.

The credential is the session's `storage_state` (cookies), not a specific
profile folder — corrected 2026-10-01, see seller_new.py's module docstring
for the live test that disproved the earlier "profile folder is the real
credential" finding. This opens a FRESH, non-persistent browser context each
run, seeded with `storage_state`, rather than reopening one specific disk
folder — so this no longer requires running from one specific machine.

SCOPE, as of 2026-10-01: only two tables are actively scraped —
BlinkitSellerHubSalesByProductRO (rolling-window item totals + Blinkit's own
computed "top selling"/transition signals, not reconstructable elsewhere) and
BlinkitSellerHubSalesOrderRO (order-level — the real item x city x day grain,
from a DIFFERENT endpoint family: the "Download sales sheet" REPORT export,
not a chart endpoint). The daily/city/category chart-histogram sweep that
used to live here (metrics/filters endpoints, per-city and per-category
fetch loops) has been removed: day, city-day and category-day totals are all
a `SUM(...) GROUP BY order_date` / `supply_city` / `business_category` on
BlinkitSellerHubSalesOrderRO, so scraping them separately only ever added
request volume, not new information. Their tables were dropped in migration
4b7e2c91d0a3.
"""
import io
import json

import httpx
from openpyxl import load_workbook
from playwright.async_api import async_playwright

from platform_auth.marketplaces.blinkit import endpoints as auth_ep
from scraper.platforms.blinkit.dashboard_data.seller_hub import endpoints as ep
from app.utils.logger import logger

_NAV_TIMEOUT_MS = 30_000
_REPORT_POLL_INTERVAL_MS = 2_000
_REPORT_POLL_ATTEMPTS = 20  # 40s — the live report took ~8-10s; generous margin
_LAUNCH_ARGS = ["--no-sandbox", "--disable-blink-features=AutomationControlled"]
_VIEWPORT = {"width": 1280, "height": 900}
# Static value from the SPA's own JS bundle, captured live alongside a real
# request (2026-10-01) — not session- or account-specific, unlike
# access_token/x-gr-seller-id below, which ARE and must come from cookies.
_X_API_KEY = "0d0b54c3-8d3a-48fc-a433-1648b22e7e8d"


DEFAULT_WINDOW = "Last 30 days"


class SessionDead(RuntimeError):
    """The stored session was logged out mid-run, after the auth probe passed.

    Seen live 2026-10-06 (Sereko): the probe reached /dashboard at 12:30:54 and
    the scrape was bounced to the login page 8 s later. The CLI catches this,
    forces a fresh login and retries once."""


async def scrape_sales(email: str, storage_state: dict, time_range_filter: str = DEFAULT_WINDOW) -> dict:
    """Fetch the item-grain rolling-window totals (products/performance) and
    the order-level report (reports/download), both via direct in-page
    fetch() — see module docstring for scope and why that's reliable now,
    not UI-driven.

    `storage_state` is the AuthSession's own cookie set (from
    `platform_auth.store.load`/`auth_service.ensure`) — see seller_new.py's
    module docstring for why this is the real credential, seeded into a
    fresh browser context rather than reopening one specific disk folder.

    `time_range_filter` must be one of the account's own valid presets ("Last
    7 days", "Last 30 days", or a named month like "September 2026" — all
    verified to return real day-grain data, not coarser). Defaults to
    `DEFAULT_WINDOW` ("Last 30 days") rather than the dashboard's own
    7-day default, so a normal scheduled run self-heals any gap under a
    month instead of losing missed days for good.
    """
    if not storage_state or not storage_state.get("cookies"):
        raise RuntimeError(
            f"No storage_state cookies for {email} — session may never have "
            "completed a real login. Check `cli auth probe blinkit_seller_new -t <uuid>`."
        )

    orders: list[dict] = []

    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=True, args=_LAUNCH_ARGS)
        context = await browser.new_context(
            storage_state=storage_state,
            viewport=_VIEWPORT,
            user_agent=auth_ep.USER_AGENT,
        )
        try:
            page = context.pages[0] if context.pages else await context.new_page()

            # Any authenticated page load is enough to clear Cloudflare and
            # establish the fetch() context — the deep Sales Performance
            # route isn't needed now that nothing here reads the SPA's own
            # auto-fired calls, only fires its own fetch()es.
            await page.goto(
                f"{ep.BASE_URL}/dashboard/home", wait_until="networkidle", timeout=_NAV_TIMEOUT_MS
            )
            if "/dashboard" not in page.url:
                raise SessionDead(
                    f"Not on a dashboard route — session may be dead. Landed on {page.url}. "
                    "Check `cli auth probe blinkit_seller_new -t <uuid>` before re-logging in."
                )
            await page.wait_for_timeout(2000)

            headers = await _auth_headers(context)

            by_product = await _fetch_json(
                page, f"/{ep.SALES_BY_PRODUCT_PATH}",
                {"filters": {"time_range_filter": time_range_filter}, "order_by": "-sales", "page": 0},
                headers,
            )

            try:
                orders = await fetch_sales_order_report(page, headers, time_range_filter)
            except Exception as e:                                    # noqa: BLE001
                logger.warning(f"Seller-hub sales order report failed, skipping: {e}")
        finally:
            await context.close()
            await browser.close()

    products = (by_product.get("data") or {}).get("product_sales_performance") or []
    logger.info(
        f"Seller-hub sales scraped — window:{time_range_filter!r} "
        f"products:{len(products)} order_rows:{len(orders)}"
    )

    return {
        "products": products,
        "window_label": time_range_filter,
        "orders": orders,
    }


async def _auth_headers(context) -> dict:
    """The 4 headers a hand-built fetch() needs beyond the auto-sent cookies
    — see module docstring. access_token/seller_id are read live off this
    run's own cookies, never hardcoded (they rotate per login)."""
    cookies = {c["name"]: c["value"] for c in await context.cookies()}
    access_token = cookies.get("access_token")
    seller_id = cookies.get("seller_id")
    if not access_token or not seller_id:
        raise SessionDead(
            "Missing access_token/seller_id cookie — session may be dead. "
            "Check `cli auth probe blinkit_seller_new -t <uuid>` before re-logging in."
        )
    return {
        "access_token": access_token,
        "app_client": "seller-dashboard-web",
        "x-api-key": _X_API_KEY,
        "x-gr-seller-id": seller_id,
    }


async def _fetch_json(
    page, path: str, body: dict | None, headers: dict, method: str = "POST"
) -> dict:
    """POST/GET a seller-hub endpoint via in-page fetch() — runs inside the
    page so it shares its origin, cookies and TLS/connection fingerprint
    (the thing Cloudflare actually gates on; see module docstring)."""
    result = await page.evaluate(
        """async ({path, body, headers, method}) => {
            const opts = {
                method,
                headers: {"Content-Type": "application/json", ...headers},
                credentials: "include",
            };
            if (method !== "GET") opts.body = JSON.stringify(body);
            const resp = await fetch(path, opts);
            const text = await resp.text();
            return {status: resp.status, text};
        }""",
        {"path": path, "body": body, "headers": headers, "method": method},
    )
    if result["status"] != 200:
        raise RuntimeError(f"{path} returned {result['status']}: {result['text'][:200]}")
    return json.loads(result["text"])


async def fetch_sales_order_report(page, headers: dict, time_range_filter: str) -> list[dict]:
    """Request, poll, and download the order-level "Download sales sheet"
    report — a genuinely different endpoint family from the chart endpoints
    (POST to start generating, poll until ready, then a plain file
    download), not UI-driven. Returns one dict per Excel data row, keyed by
    the sheet's own column headers (e.g. "Order Id", "Order Date",
    "Item Id", ...) — parser.py does the field-name/type mapping.

    Verified live (Sereko, 2026-10-01): requesting "Last 30 days" returned a
    September-2026 report whose total matched the September total computed
    independently from the (now-retired) daily chart endpoint to the rupee
    (Rs 3,36,604) — confirmed three separate times. This is the real item x
    city x day grain.
    """
    init = await _fetch_json(
        page,
        f"/{ep.REPORTS_DOWNLOAD_PATH}",
        {"report_type": "SALES_SUMMARY", "params": {"time_range_filter": time_range_filter}},
        headers,
    )
    poll_key = (init.get("data") or {}).get("poll_key")
    if not poll_key:
        raise RuntimeError(f"reports/download didn't return a poll_key: {init}")

    document_url = None
    for _ in range(_REPORT_POLL_ATTEMPTS):
        await page.wait_for_timeout(_REPORT_POLL_INTERVAL_MS)
        poll = await _fetch_json(
            page, f"/{ep.REPORTS_POLL_PATH}?poll_key={poll_key}", None, headers, method="GET",
        )
        data = poll.get("data") or {}
        if data.get("document_url"):
            document_url = data["document_url"]
            break
        if data.get("status") not in ("PENDING", None):
            raise RuntimeError(f"reports/poll returned an unexpected status: {data}")
    if not document_url:
        raise RuntimeError(
            f"Report for {time_range_filter!r} never finished after "
            f"{_REPORT_POLL_ATTEMPTS * _REPORT_POLL_INTERVAL_MS / 1000:.0f}s of polling"
        )

    # A plain, direct download — presigned S3 links aren't behind Blinkit's
    # Cloudflare at all, so no browser/fetch() dance needed for this step.
    resp = httpx.get(document_url, timeout=30)
    resp.raise_for_status()

    wb = load_workbook(io.BytesIO(resp.content), data_only=True)
    ws = wb["Sales Summary"]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header, *data_rows = rows
    return [dict(zip(header, row)) for row in data_rows if any(c is not None for c in row)]
